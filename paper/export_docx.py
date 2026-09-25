#!/usr/bin/env python3
"""Export manuscript.md to an editable Word document, manuscript.docx.

The layout mirrors manuscript.typ: US Letter, 10-point Times New Roman, a
full-width title block, and a two-column body. Requires python-docx
(`pip install python-docx`). Neither manuscript.md nor manuscript.typ is modified.
"""
from pathlib import Path
import re

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    raise SystemExit("Install python-docx (pip install python-docx), then rerun this script.")

ROOT = Path(__file__).resolve().parent
FONT = "Times New Roman"
MARGIN_X, GUTTER = 0.65, 0.25
COLUMN = (8.5 - 2*MARGIN_X - GUTTER) / 2
ROMAN = {"1":"I", "2":"II", "3":"III", "4":"IV", "5":"V", "6":"VI"}
# Same ASCII equation layout and definitions as export_typst.py.
EQUATIONS = {
    1: (["v_i(t) = 1{m_i(t) > q_0.45^s(t)}."],
        "Here, m_i(t) is the 1 s sound mean, and q_p^x(t) is the rolling p-quantile of signal x_i over 60 s. The indicator 1{condition} equals 1 when the condition holds and 0 otherwise."),
    2: (["z_i^x = (x_i - q_0.25^x) / max(IQR_i^x, 1)."],
        "Here, x is either sound s or acceleration a; ^x identifies the signal rather than an exponent. Time arguments are omitted in (2), and IQR_i^x = q_0.75^x - q_0.25^x."),
    3: (["c_i(t) = z_i^s(t) + 1.2 * clip(z_i^a(t), 0, 2.5)."], None),
    4: (["H = -sum_i(p_i * ln(p_i)) / ln(n)."], None),
    5: (["G = sum_i sum_j |d_i - d_j| / (2n sum_i d_i)."], None),
}
TABLE_WIDTHS = {"Number": (1.0, 3.3, 0.85), "Session-held-out": (1.6, 0.7, 1.1), 2: (1.2, 3.5), 3: (2, 1, 1)}


def hyperlink(paragraph, url, text, size=None):
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x93)
    if size:
        run.font.size = size
    link.append(run._r)
    paragraph._p.append(link)


def inline(paragraph, source, size=None, citations=True):
    tokens = re.compile(r'(`[^`]+`|\*\*.+?\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))')
    for i, chunk in enumerate(tokens.split(source)):
        if not chunk:
            continue
        if i % 2 == 0:
            run = paragraph.add_run(chunk)
        elif chunk.startswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(7.5)
            continue
        elif chunk.startswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', chunk).groups()
            if citations and label.isdigit():
                # Citations follow the preceding sentence, as in the IEEE style.
                if paragraph.runs and paragraph.runs[-1].text.endswith(" "):
                    paragraph.runs[-1].text = paragraph.runs[-1].text.rstrip()
                run = paragraph.add_run(" ["+label+"]")
            else:
                hyperlink(paragraph, url, label, size)
                continue
        if size:
            run.font.size = size


def set_columns(section, count):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(GUTTER*1440)))


def plain(doc, align=None, indent=True, size=None, before=0, after=0):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    if align is not None:
        p.alignment = align
    if not indent:
        fmt.first_line_indent = Pt(0)
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    return p


def caption(doc, text):
    p = plain(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False, before=3, after=6)
    label, body = re.match(r'\*((?:Figure|Table) \d+\.) (.*)\*$', text).groups()
    p.add_run(label+" ").font.size = Pt(8)
    inline(p, body, Pt(8))
    return p


def table(doc, block):
    rows = [[c.strip() for c in r.strip().strip("|").split("|")]
            for r in block.splitlines()]
    del rows[1]
    key = next((k for k in TABLE_WIDTHS if isinstance(k, str) and k in "|".join(rows[0])), len(rows[0]))
    widths = TABLE_WIDTHS[key]
    scale = COLUMN / sum(widths)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for c, column in enumerate(t.columns):
        column.width = Inches(widths[c]*scale)
    for r, row in enumerate(rows):
        if r == 0:
            header = OxmlElement("w:tblHeader")
            t.rows[0]._tr.get_or_add_trPr().append(header)
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Inches(widths[c]*scale)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p, "**"+text+"**" if r == 0 else text, Pt(8))
    plain(doc, indent=False, after=4).paragraph_format.line_spacing = Pt(4)


def equation(doc, number):
    lines, note = EQUATIONS[number]
    for k, line in enumerate(lines):
        p = plain(doc, indent=False, before=5 if k == 0 else 0,
                  after=5 if k == len(lines)-1 else 1)
        p.paragraph_format.keep_with_next = k < len(lines)-1
        stops = p.paragraph_format.tab_stops
        stops.add_tab_stop(Inches(COLUMN/2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(Inches(COLUMN), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run("\t"+line+("\t("+str(number)+")" if k == len(lines)-1 else ""))
        run.font.size = Pt(9)
    if note:
        inline(plain(doc, WD_ALIGN_PARAGRAPH.JUSTIFY), note)


def setup(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(MARGIN_X)
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    fmt = normal.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Pt(10)
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for level, italic in ((1, False), (2, True)):
        style = doc.styles["Heading %d" % level]
        style.font.name = FONT
        style.font.size = Pt(10)
        style.font.bold = False
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        style.font.all_caps = level == 1
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Pt(0)
        pf.space_before, pf.space_after = (Pt(11), Pt(6)) if level == 1 else (Pt(8), Pt(4))
        pf.keep_with_next = True


def main():
    blocks = (ROOT/"manuscript.md").read_text().strip().split("\n\n")
    doc = Document()
    setup(doc)
    i = 0
    while i < len(blocks):
        block = blocks[i].strip()
        if block.startswith(("**Manuscript draft", "**[AUTHOR QUERY:")):
            pass
        elif block.startswith("# "):
            title, authors, affiliation = block[2:], blocks[i+1], blocks[i+2]
            doc.core_properties.title = title
            doc.core_properties.author = authors.replace(", and ", ", ").replace(" and ", ", ")
            p = plain(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=8)
            p.add_run(title).font.size = Pt(22)
            p.paragraph_format.line_spacing = 1.0
            plain(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False).add_run(authors).font.size = Pt(11)
            plain(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=18).add_run(affiliation)
            set_columns(doc.sections[0], 1)
            set_columns(doc.add_section(WD_SECTION.CONTINUOUS), 2)
            i += 2
        elif block == "## Abstract":
            p = plain(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.add_run("Abstract: ").bold = True
            inline(p, blocks[i+1])
            i += 1
        elif block == "## References":
            doc.add_heading("References", 1)
            for ref in blocks[i+1:]:
                number, text = re.match(r'\[(\d+)\] (.*)', ref).groups()
                p = plain(doc, WD_ALIGN_PARAGRAPH.LEFT, after=2)
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-18)
                p.add_run("["+number+"]\t").font.size = Pt(8)
                p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
                inline(p, text, Pt(8), citations=False)
            break
        elif block.startswith("##"):
            hashes, title = re.match(r'^(#+) (.*)$', block).groups()
            if len(hashes) == 2:
                title = re.sub(r'^(\d+)\.', lambda m: ROMAN[m[1]]+'.', title)
            else:
                title = re.sub(r'^\d+\.(\d+)\.', lambda m: chr(64+int(m[1]))+'.', title)
            doc.add_heading(title, len(hashes)-1)
        elif block.startswith("**Equation"):
            equation(doc, int(re.match(r'\*\*Equation (\d+)', block)[1]))
        elif block.startswith("*Table "):
            caption(doc, block).paragraph_format.keep_with_next = True
        elif block.startswith("| "):
            table(doc, block)
        elif block.startswith("!["):
            alt, path = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', block).groups()
            p = plain(doc, WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=4)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ROOT/path), width=Inches(COLUMN))
            doc.inline_shapes[-1]._inline.docPr.set("descr", alt)
            caption(doc, blocks[i+1])
            i += 1
        else:
            block = re.sub(r'\*\*\[AUTHOR QUERY:.*?\]\*\*', '', block).strip()
            inline(plain(doc, WD_ALIGN_PARAGRAPH.JUSTIFY), block)
        i += 1
    doc.save(ROOT/"manuscript.docx")
    print(f"Wrote {ROOT/'manuscript.docx'}")


if __name__ == "__main__":
    main()
